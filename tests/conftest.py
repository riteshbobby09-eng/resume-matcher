"""
Shared test fixtures for Resume Matcher.

Provides:
- Temporary directories
- Sample text data
- Mock documents
- Test configuration
- Synthetic PDF/DOCX fixture generators
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument

from resume_matcher.config import AppConfig, load_config
from resume_matcher.domain.enums import DocumentType, FileFormat, SectionType
from resume_matcher.domain.models import (
    Block,
    DocumentMeta,
    IndexedLine,
    MatchEvidence,
    ProcessedDocument,
    Sentence,
)


# ---------------------------------------------------------------------------
# Configuration Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="session")
def test_config() -> AppConfig:
    """Load test configuration."""
    test_config_path = Path("config/test.yaml")
    if test_config_path.exists():
        return load_config(config_path="config/default.yaml", override_path=str(test_config_path))
    return load_config(config_path="config/default.yaml")


# ---------------------------------------------------------------------------
# Temporary Directory Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def tmp_dir():
    """Create a temporary directory for test files."""
    with tempfile.TemporaryDirectory(prefix="resume_matcher_test_") as d:
        yield Path(d)


# ---------------------------------------------------------------------------
# Sample Text Data
# ---------------------------------------------------------------------------
SAMPLE_RESUME_TEXT = """John Doe
Email: john.doe@email.com | Phone: +1-555-123-4567

PROFESSIONAL SUMMARY
Experienced software engineer with 5+ years of experience in Python, Java, and cloud technologies.
Proven track record of delivering scalable microservices and data pipelines.

SKILLS
Python, Java, JavaScript, SQL, Docker, Kubernetes, AWS, GCP
Machine Learning, TensorFlow, PyTorch, Scikit-learn
Git, CI/CD, Jenkins, GitHub Actions
PostgreSQL, MongoDB, Redis, Elasticsearch

WORK EXPERIENCE
Senior Software Engineer | TechCorp Inc. | 2021 - Present
• Led development of a microservices platform serving 10M+ users
• Designed and implemented real-time data processing pipeline using Apache Kafka
• Reduced deployment time by 60% through CI/CD automation
• Mentored team of 5 junior developers

Software Engineer | DataSoft Solutions | 2019 - 2021
• Built RESTful APIs using Python Flask and FastAPI
• Implemented machine learning models for customer churn prediction
• Worked with cross-functional teams in Agile environment
• Achieved 99.9% uptime for production services

EDUCATION
Bachelor of Technology in Computer Science
Indian Institute of Technology, Delhi | 2015 - 2019
CGPA: 8.7/10.0

CERTIFICATIONS
AWS Solutions Architect - Associate
Google Cloud Professional Data Engineer

PROJECTS
• Open-source contribution to pandas library
• Personal project: Real-time stock market dashboard using React and Python
"""

SAMPLE_JD_TEXT = """Job Title: Senior Software Engineer

About Us
We are a leading technology company building innovative solutions for enterprise customers.

Requirements
Must have 5+ years of experience in software development.
Required: Strong proficiency in Python and at least one other programming language.
Experience with cloud platforms (AWS/GCP/Azure) is essential.
Must have experience with microservices architecture and containerization.
Knowledge of machine learning concepts is required.

Responsibilities
• Design and implement scalable backend services
• Lead code reviews and mentor junior team members
• Collaborate with product teams to define technical requirements
• Drive continuous improvement in development practices

Skills Required
Python, Java or Go, SQL, Docker, Kubernetes
Cloud platforms: AWS or GCP
CI/CD tools: Jenkins, GitHub Actions
Database technologies: PostgreSQL, MongoDB

Education
Bachelor's degree in Computer Science or related field.
Master's degree preferred.

Nice to Have
• Experience with real-time data processing
• Contributions to open-source projects
• Knowledge of system design and distributed systems
"""


@pytest.fixture
def sample_resume_lines() -> list[IndexedLine]:
    """Create sample IndexedLines from resume text."""
    lines = []
    for i, text in enumerate(SAMPLE_RESUME_TEXT.strip().split("\n"), start=1):
        lines.append(
            IndexedLine(
                global_line_number=i,
                text=text,
                page_number=1,
                local_line_number=i,
                source_filename="test_resume.pdf",
                document_id="resume_001",
                is_empty=(text.strip() == ""),
            )
        )
    return lines


@pytest.fixture
def sample_jd_lines() -> list[IndexedLine]:
    """Create sample IndexedLines from JD text."""
    lines = []
    for i, text in enumerate(SAMPLE_JD_TEXT.strip().split("\n"), start=1):
        lines.append(
            IndexedLine(
                global_line_number=i,
                text=text,
                page_number=1,
                local_line_number=i,
                source_filename="test_jd.pdf",
                document_id="jd_001",
                is_empty=(text.strip() == ""),
            )
        )
    return lines


# ---------------------------------------------------------------------------
# File Fixtures (create actual PDF/DOCX files)
# ---------------------------------------------------------------------------
@pytest.fixture
def sample_pdf(tmp_dir) -> Path:
    """Create a sample PDF file with resume content."""
    path = tmp_dir / "test_resume.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    text_point = pymupdf.Point(72, 72)
    for line in SAMPLE_RESUME_TEXT.strip().split("\n")[:30]:
        page.insert_text(text_point, line, fontsize=10)
        text_point.y += 14
        if text_point.y > 750:
            page = doc.new_page()
            text_point = pymupdf.Point(72, 72)
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def sample_docx(tmp_dir) -> Path:
    """Create a sample DOCX file with resume content."""
    path = tmp_dir / "test_resume.docx"
    doc = DocxDocument()
    for line in SAMPLE_RESUME_TEXT.strip().split("\n"):
        doc.add_paragraph(line)
    doc.save(str(path))
    return path


@pytest.fixture
def sample_jd_pdf(tmp_dir) -> Path:
    """Create a sample PDF file with JD content."""
    path = tmp_dir / "test_jd.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    text_point = pymupdf.Point(72, 72)
    for line in SAMPLE_JD_TEXT.strip().split("\n")[:30]:
        page.insert_text(text_point, line, fontsize=10)
        text_point.y += 14
        if text_point.y > 750:
            page = doc.new_page()
            text_point = pymupdf.Point(72, 72)
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def empty_pdf(tmp_dir) -> Path:
    """Create an empty PDF file (no text)."""
    path = tmp_dir / "empty.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def empty_docx(tmp_dir) -> Path:
    """Create an empty DOCX file."""
    path = tmp_dir / "empty.docx"
    doc = DocxDocument()
    doc.save(str(path))
    return path


@pytest.fixture
def docx_with_table(tmp_dir) -> Path:
    """Create a DOCX file containing a table."""
    path = tmp_dir / "table_resume.docx"
    doc = DocxDocument()
    doc.add_paragraph("SKILLS")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Language"
    table.cell(0, 1).text = "Level"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "Expert"
    table.cell(2, 0).text = "Java"
    table.cell(2, 1).text = "Intermediate"
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Senior Developer at TechCorp, 2020 - Present")
    doc.save(str(path))
    return path
